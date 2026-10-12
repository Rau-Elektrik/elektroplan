"""Excel-Vorlage „Verteilerdokumentation Stromkreisliste Messprotokoll“ mit den eigenen Firmendaten.

Wird bei der Einrichtung erzeugt (einrichtung.py) und landet im Konfigurationsordner
(~/.elektroplan/vorlagen/). elektroplan.py --stromkreisliste füllt sie anschließend aus.

Konvention: gelbe Zellen = Eingabe, weiße Zellen = Formel (nicht überschreiben).
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

from . import firma as _firma

FIRMA: dict = {}
LOGO_XL: Path | None = None
LOGO_RATIO = 1.0

# ---------------------------------------------------------------- Gestaltung
GRAPHIT, GRAU, GRAU_HELL = "1C1C1C", "6E6A66", "A9A39C"
KUPFER = "B87333"          # Akzentfarbe, wird aus firma.json überschrieben
LINIE, FLAECHE, EINGABE = "D6D1CB", "F5F1EC", "FFF5C8"
SCHRIFT = "Arial"

DUENN = Side(style="thin", color=LINIE)
RAHMEN = Border(left=DUENN, right=DUENN, top=DUENN, bottom=DUENN)
KUPFER_LINIE = Side(style="medium", color=KUPFER)
STRICH = Side(style="thin", color=GRAPHIT)

FMT_DATUM = "DD.MM.YYYY"


def font(size=10, bold=False, color=GRAPHIT, italic=False):
    return Font(name=SCHRIFT, size=size, bold=bold, color=color, italic=italic)


def fill(farbe):
    return PatternFill("solid", fgColor=farbe)


def zelle(ws, ref, wert=None, *, size=10, bold=False, color=GRAPHIT, italic=False, bg=None,
          h="left", v="center", wrap=False, rahmen=False, fmt=None):
    c = ws[ref]
    if wert is not None:
        c.value = wert
    c.font = font(size, bold, color, italic)
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if bg:
        c.fill = fill(bg)
    if rahmen:
        c.border = RAHMEN
    if fmt:
        c.number_format = fmt
    return c


def bereich(ws, rng, **kw):
    """Formatiert alle Zellen eines Bereichs gleich (z. B. vor dem Verbinden)."""
    if ":" not in rng:
        zelle(ws, rng, **kw)
        return
    for row in ws[rng]:
        for c in row:
            zelle(ws, c.coordinate, **kw)


def verbinden(ws, rng, wert=None, **kw):
    bereich(ws, rng, **kw)
    ws.merge_cells(rng)
    if wert is not None:
        ws[rng.split(":")[0]].value = wert


def eingabe(ws, rng, fmt=None, h="left", wrap=False, merge=False, wert=None):
    if merge:
        verbinden(ws, rng, wert, bg=EINGABE, rahmen=True, fmt=fmt, h=h, wrap=wrap)
    else:
        bereich(ws, rng, bg=EINGABE, rahmen=True, fmt=fmt, h=h, wrap=wrap)
        if wert is not None:
            ws[rng.split(":")[0]].value = wert


def beschriftung(ws, ref, text):
    zelle(ws, ref, text, size=8, color=GRAU, bg=FLAECHE, rahmen=True, wrap=True)


def formelfeld(ws, rng, formel, fmt=None, h="left", merge=False, bold=False):
    if merge:
        verbinden(ws, rng, formel, rahmen=True, fmt=fmt, h=h, bold=bold)
    else:
        zelle(ws, rng, formel, rahmen=True, fmt=fmt, h=h, bold=bold)


def abschnitt(ws, zeile, text, von="A", bis="D"):
    verbinden(ws, f"{von}{zeile}:{bis}{zeile}", text.upper(), size=9, bold=True, color=GRAPHIT)
    for col in range(ws[f"{von}{zeile}"].column, ws[f"{bis}{zeile}"].column + 1):
        ws.cell(zeile, col).border = Border(bottom=KUPFER_LINIE)
    ws.row_dimensions[zeile].height = 20


def kopfzeile(ws, zeile, spalten, hoehe=30):
    """Tabellenkopf in Graphit. spalten: Liste von (Spaltenbuchstabe, Text)."""
    for col, text in spalten:
        zelle(ws, f"{col}{zeile}", text, size=8, bold=True, color="FFFFFF", bg=GRAPHIT, h="center", wrap=True,
              rahmen=True)
    ws.row_dimensions[zeile].height = hoehe


def breiten(ws, werte):
    for col, w in werte.items():
        ws.column_dimensions[col].width = w


def spalte_px(ws, col):
    w = ws.column_dimensions[col].width or 8.43
    return int(w * 7 + 0.5)


def logo(ws, bis_spalte, breite_px=130, zeile=0, oben_px=4):
    """Setzt das Logo rechtsbündig an das Ende von bis_spalte."""
    if LOGO_XL is None:
        return
    from openpyxl.utils import column_index_from_string
    ende = column_index_from_string(bis_spalte)
    rest = breite_px
    col = ende
    while col >= 1:
        px = spalte_px(ws, get_column_letter(col))
        if rest <= px:
            off = px - rest
            break
        rest -= px
        col -= 1
    img = XLImage(str(LOGO_XL))
    hoehe_px = round(breite_px / LOGO_RATIO)
    img.width, img.height = breite_px, hoehe_px
    marker = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(off), row=zeile, rowOff=pixels_to_EMU(oben_px))
    img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(breite_px), pixels_to_EMU(hoehe_px)))
    ws.add_image(img)


def titel(ws, text, unter, bis_spalte, zeile_kupfer=3):
    zelle(ws, "A1", text, size=16, bold=True)
    zelle(ws, "A2", unter, size=9, color=GRAU)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[zeile_kupfer].height = 8
    from openpyxl.utils import column_index_from_string
    for col in range(1, column_index_from_string(bis_spalte) + 1):
        ws.cell(zeile_kupfer, col).border = Border(bottom=KUPFER_LINIE)
    logo(ws, bis_spalte)


def liste(ws, rng, werte, fehler=False):
    dv = DataValidation(type="list", formula1='"' + ",".join(werte) + '"', allow_blank=True,
                        showErrorMessage=fehler)
    ws.add_data_validation(dv)
    dv.add(rng)


def drucken(ws, bereich_, quer=False, titelzeilen=None, eine_seite=False, skalierung=None):
    ws.print_area = bereich_
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if quer else "portrait"
    if skalierung:
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = skalierung
    else:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1 if eine_seite else 0
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.55, bottom=0.65, header=0.25, footer=0.3)
    ws.print_options.horizontalCentered = True
    if titelzeilen:
        ws.print_title_rows = titelzeilen
    ws.oddFooter.left.text = _verbinde(FIRMA.get("name"), _person(), FIRMA.get("telefon"), FIRMA.get("mail"))
    ws.oddFooter.left.size = 7
    ws.oddFooter.left.font = "Arial,Regular"
    ws.oddFooter.left.color = GRAU
    ws.oddFooter.right.text = "Seite &P von &N"
    ws.oddFooter.right.size = 7
    ws.oddFooter.right.font = "Arial,Regular"
    ws.oddFooter.right.color = GRAU
    ws.sheet_view.showGridLines = False


def hinweis(ws, rng, text, hoehe=28):
    verbinden(ws, rng, text, size=8, italic=True, color=GRAU, wrap=True, v="top")
    ws.row_dimensions[int("".join(ch for ch in rng.split(":")[0] if ch.isdigit()))].height = hoehe


def ergebnis_farben(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"n.i.O."'],
                                                  font=Font(name=SCHRIFT, bold=True, color="C00000"),
                                                  fill=fill("FDE2E2")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"i.O."'],
                                                  font=Font(name=SCHRIFT, bold=True, color="2E7D32"),
                                                  fill=fill("E6F4EA")))


# ============================================================================
#  1  Verteilerdokumentation
# ============================================================================
SK = "Stromkreisliste"
MP = "Messprotokoll"
PB = "'Prüfbericht'"
ERSTE, LETZTE = 7, 56  # Datenzeilen Stromkreisliste / Messprotokoll


def blatt_anleitung(wb):
    ws = wb.active
    ws.title = "Anleitung"
    breiten(ws, {"A": 4, "B": 96})
    titel(ws, "Verteilerdokumentation", "So arbeitest du mit dieser Vorlage", "B")
    zeilen = [
        ("", "Ablauf", True),
        ("1", "Vorlage per Doppelklick öffnen – Excel legt eine neue Datei an. Gleich im Kundenordner speichern, "
              "z. B. als „JJJJ-MM-TT Verteilerdokumentation <Verteiler>.xlsx“.", False),
        ("2", "Blatt „Anlage“: Kunde, Verteiler, Netz und Erdung ausfüllen. Diese Angaben erscheinen automatisch auf allen anderen Blättern.", False),
        ("3", "Blatt „Stromkreisliste“: eine Zeile je Schutzorgan (FI, LS, Hauptschalter …). Reihe und Platz = Position im Verteiler "
              "(Reihe 1 oben, Platz 1 links, in Teilungseinheiten). Die Spalte „Beschriftung kurz“ wird zum Beschriftungsstreifen.", False),
        ("4", "Blatt „Beschriftung“: füllt sich selbst. Mit 100 % drucken (keine Skalierung), ein Feld = 1 TE = 17,5 mm. "
              "Am Probedruck mit dem Lineal nachmessen, ausschneiden, in die Beschriftungsleiste des Verteilers schieben.", False),
        ("5", "Blatt „Messprotokoll“: Messwerte je Stromkreis eintragen. Ik, Zs max und das Ergebnis rechnen sich selbst. "
              "Die Grenzwerte stehen oben im Blatt und lassen sich anpassen.", False),
        ("6", "Blatt „Prüfbericht“: Besichtigen und Erproben abhaken, Mängel eintragen, Gesamtergebnis wählen, unterschreiben.", False),
        ("7", "Für den Kunden: Stromkreisliste, Prüfbericht und Messprotokoll als PDF speichern (Datei › Exportieren) und mit der "
              "Kundendokumentation übergeben. Eine Kopie der Stromkreisliste gehört in den Verteiler.", False),
        ("", "", False),
        ("", "Farben", True),
        ("", "Gelbe Zellen = deine Eingabe.  Weiße Zellen = Formel, nicht überschreiben.  Zeilen 7–9 der Stromkreisliste sind "
             "Beispiele – überschreiben oder löschen.", False),
        ("", "Leere Zeilen beim Drucken ausblenden: in der Kopfzeile den Filter bei „BMK“ öffnen und „(Leere)“ abwählen.", False),
        ("", "", False),
        ("", "Wichtig", True),
        ("", "Die automatische Bewertung ist eine Hilfe, kein Ersatz für die Beurteilung durch die Elektrofachkraft. "
             "Maßgeblich sind DIN VDE 0100-600, DIN VDE 0105-100 und die Angaben der Gerätehersteller.", False),
    ]
    r = 5
    for nr, text, kopf in zeilen:
        if kopf:
            abschnitt(ws, r, text, "A", "B")
        else:
            zelle(ws, f"A{r}", nr, bold=True, color=KUPFER, v="top", h="center")
            zelle(ws, f"B{r}", text, wrap=True, v="top", size=10)
            ws.row_dimensions[r].height = 15 if len(text) < 100 else 28 if len(text) < 200 else 41
        r += 1
    drucken(ws, f"A1:B{r}", eine_seite=True)


def blatt_anlage(wb):
    ws = wb.create_sheet("Anlage")
    breiten(ws, {"A": 27, "B": 42, "C": 22, "D": 30})
    titel(ws, "Verteilerdokumentation", "Anlagen- und Verteilerdaten", "D")
    hinweis(ws, "A4:D4", "Gelbe Felder ausfüllen – die Angaben erscheinen automatisch auf Stromkreisliste, Beschriftung, "
                         "Prüfbericht und Messprotokoll.", 24)

    def feld(r, lab, ref, lab2=None, ref2=None, fmt=None, fmt2=None, span=False, wert=None, wert2=None):
        beschriftung(ws, f"A{r}", lab)
        if span:
            eingabe(ws, f"B{r}:D{r}", merge=True, fmt=fmt, wert=wert)
        else:
            eingabe(ws, f"B{r}", fmt=fmt, wert=wert)
            if lab2:
                beschriftung(ws, f"C{r}", lab2)
                eingabe(ws, f"D{r}", fmt=fmt2, wert=wert2)
        ws.row_dimensions[r].height = 20

    abschnitt(ws, 6, "Kunde und Objekt")
    feld(7, "Auftraggeber / Betreiber", "B7", "Kunden-Nr.", "D7")
    feld(8, "Anlagenanschrift", "B8", span=True)
    feld(9, "Auftrags-Nr.", "B9", "Datum", "D9", fmt2=FMT_DATUM)
    abschnitt(ws, 11, "Verteiler")
    feld(12, "Bezeichnung (z. B. UV EG)", "B12", "Standort", "D12")
    feld(13, "Hersteller / Typ", "B13", "Reihen × TE", "D13")
    feld(14, "Zuleitung (Typ, Querschnitt)", "B14", "Vorsicherung", "D14")
    abschnitt(ws, 16, "Netz und Erdung")
    feld(17, "Netzbetreiber", "B17", "Zählernummer", "D17")
    feld(18, "Netzsystem", "B18", "Nennspannung", "D18", wert2="230/400 V, 50 Hz")
    feld(19, "Hausanschlusssicherung", "B19", "Hauptpotentialausgleich", "D19")
    feld(20, "Überspannungsschutz", "B20", "Erdungswiderstand RA (Ω)", "D20", fmt2="0.00")
    feld(21, "Erdungsanlage", "B21")
    liste(ws, "B18", ["TN-C-S", "TN-S", "TT", "TN-C"])
    liste(ws, "D19", ["vorhanden", "nicht vorhanden"])
    liste(ws, "B20", ["Typ 1+2", "Typ 2", "Typ 3", "nicht vorhanden"])
    liste(ws, "B21", ["Fundamenterder", "Ringerder", "Tiefenerder", "unbekannt"])

    abschnitt(ws, 23, "Errichter und Prüfer")
    beschriftung(ws, "A24", "Errichter")
    verbinden(ws, "B24:D24", _verbinde(FIRMA.get("name"), _person(), FIRMA.get("strasse"), FIRMA.get("ort")),
              rahmen=True)
    beschriftung(ws, "A25", "Installateurverzeichnis")
    zelle(ws, "B25", FIRMA.get("installateurverzeichnis") or None, rahmen=True)
    beschriftung(ws, "C25", "HWK-Betriebs-Nr.")
    zelle(ws, "D25", FIRMA.get("hwk_nr") or None, rahmen=True)
    beschriftung(ws, "A26", "Prüfer")
    eingabe(ws, "B26", wert=FIRMA.get("inhaber") or None)
    beschriftung(ws, "C26", "Telefon")
    zelle(ws, "D26", FIRMA.get("telefon") or None, rahmen=True)
    for r in (24, 25, 26):
        ws.row_dimensions[r].height = 20

    abschnitt(ws, 28, "Zusammenfassung (automatisch)")
    beschriftung(ws, "A29", "Stromkreise erfasst")
    formelfeld(ws, "B29", f'=COUNTIF({SK}!B{ERSTE}:B{LETZTE},"?*")', h="left")
    beschriftung(ws, "C29", "Stromkreise gemessen")
    formelfeld(ws, "D29", f'=COUNTIF({MP}!P{ERSTE}:P{LETZTE},"?*")', h="left")
    beschriftung(ws, "A30", "davon in Ordnung")
    formelfeld(ws, "B30", f'=COUNTIF({MP}!P{ERSTE}:P{LETZTE},"i.O.")', h="left")
    beschriftung(ws, "C30", "davon nicht in Ordnung")
    formelfeld(ws, "D30", f'=COUNTIF({MP}!P{ERSTE}:P{LETZTE},"n.i.O.")', h="left")
    beschriftung(ws, "A31", "Ergebnis der Messungen")
    formelfeld(ws, "B31:D31", '=IF(D30>0,"Mängel festgestellt – siehe Messprotokoll",'
                              'IF(B30>0,"Alle gemessenen Stromkreise in Ordnung","noch keine Messwerte"))',
               merge=True, bold=True)
    for r in (29, 30, 31):
        ws.row_dimensions[r].height = 20
    ws.conditional_formatting.add("B31", FormulaRule(formula=["$D$30>0"], font=Font(name=SCHRIFT, bold=True, color="C00000")))
    drucken(ws, "A1:D31", eine_seite=True)
    ws.freeze_panes = None


def blatt_stromkreisliste(wb):
    ws = wb.create_sheet(SK)
    spalten = {
        "A": ("Nr.", 4.5), "B": ("BMK", 7), "C": ("Reihe", 6), "D": ("Platz (TE)", 6.5), "E": ("Breite (TE)", 6.5),
        "F": ("Schutzorgan", 14), "G": ("Char.", 5.5), "H": ("In (A)", 6), "I": ("Pole", 6),
        "J": ("vor-geschalteter FI", 9), "K": ("IΔn (mA)", 7), "L": ("FI-Typ", 6),
        "M": ("Leitung", 10), "N": ("Adern × mm²", 9), "O": ("Länge ca. (m)", 7.5),
        "P": ("Raum / Bereich", 16), "Q": ("Verbraucher / Beschreibung", 30), "R": ("Beschriftung kurz (max. 14 Zeichen)", 17),
        "S": ("Schlüssel", 6),
    }
    breiten(ws, {k: v[1] for k, v in spalten.items()})
    titel(ws, "Stromkreisliste", f"Verteilerdokumentation · {FIRMA.get('name', '')}", "R")
    # Kopfdaten aus "Anlage"
    beschriftung(ws, "A4", "Verteiler")
    ws.merge_cells("A4:B4")
    formelfeld(ws, "C4:F4", '=IF(Anlage!B12="","",Anlage!B12&IF(Anlage!D12="",""," · "&Anlage!D12))', merge=True, bold=True)
    beschriftung(ws, "G4", "Objekt")
    ws.merge_cells("G4:I4")
    formelfeld(ws, "J4:O4", '=IF(Anlage!B8="","",Anlage!B8)', merge=True)
    beschriftung(ws, "P4", "Auftraggeber")
    formelfeld(ws, "Q4", '=IF(Anlage!B7="","",Anlage!B7)')
    formelfeld(ws, "R4", '=IF(Anlage!D9="","",Anlage!D9)', fmt=FMT_DATUM, h="center")
    ws.row_dimensions[4].height = 20
    hinweis(ws, "A5:R5", "Eine Zeile je Schutzorgan. Reihe / Platz = Position im Verteiler (Reihe 1 oben, Platz 1 links) – "
                         "daraus entsteht das Blatt „Beschriftung“. Zeilen 7–9 sind Beispiele: überschreiben oder löschen.", 18)
    kopfzeile(ws, 6, [(k, v[0]) for k, v in spalten.items() if k != "S"], hoehe=34)
    zelle(ws, "S6", "Schlüssel", size=7, color=GRAU_HELL)

    for r in range(ERSTE, LETZTE + 1):
        zelle(ws, f"A{r}", f'=IF(B{r}="","",ROW()-{ERSTE - 1})', rahmen=True, h="center", size=9, color=GRAU)
        for col in "BCDEFGHIJKLMNOPQR":
            zelle(ws, f"{col}{r}", bg=EINGABE, rahmen=True, size=9,
                  h="center" if col in "CDEGHIKLO" else "left")
        zelle(ws, f"S{r}", f'=IF(OR(C{r}="",D{r}=""),"",C{r}&"-"&D{r})', size=7, color=GRAU_HELL)
        ws.row_dimensions[r].height = 16

    beispiele = [
        ["F1", 1, 1, 4, "FI/RCD", None, 40, 4, None, 30, "A", None, None, None, "Allgemein", "FI für F2–F6", "FI 40A 30mA"],
        ["F2", 1, 5, 1, "LS", "B", 16, 1, "F1", None, None, "NYM-J", "3×2,5", 18, "Küche", "Steckdosen Arbeitsplatte", "Küche SD"],
        ["F3", 1, 6, 1, "LS", "B", 10, 1, "F1", None, None, "NYM-J", "3×1,5", 22, "Flur", "Licht Flur und Treppe", "Licht Flur"],
    ]
    for i, werte in enumerate(beispiele):
        for j, w in enumerate(werte):
            if w is not None:
                ws.cell(ERSTE + i, 2 + j).value = w

    rng = lambda c: f"{c}{ERSTE}:{c}{LETZTE}"
    liste(ws, rng("F"), ["LS", "FI/RCD", "FI/LS (RCBO)", "Hauptschalter", "NH-Sicherung", "Neozed", "SLS",
                         "Überspannungsschutz", "Sonstiges"])
    liste(ws, rng("G"), ["B", "C", "D", "K", "Z", "gG"])
    liste(ws, rng("I"), ["1", "1+N", "2", "3", "3+N", "4"])
    liste(ws, rng("K"), ["10", "30", "100", "300", "500"])
    liste(ws, rng("L"), ["A", "F", "B", "B+", "AC", "EV"])
    liste(ws, rng("M"), ["NYM-J", "NYM-O", "NYY-J", "NHXMH-J", "H07V-K", "H07RN-F", "Sonstige"])
    ws.conditional_formatting.add(rng("R"), FormulaRule(formula=[f"LEN($R{ERSTE})>14"], fill=fill("F8CBAD")))
    ws["R6"].comment = Comment("Text für den Beschriftungsstreifen. Mehr als 14 Zeichen passen nicht in 1 TE "
                               "und werden orange markiert.", _autor())
    ws["D6"].comment = Comment("Platz = erste Teilungseinheit (TE) des Geräts, von links gezählt.", _autor())
    ws.column_dimensions["S"].hidden = True
    ws.auto_filter.ref = f"A6:R{LETZTE}"
    ws.freeze_panes = "C7"
    drucken(ws, f"A1:R{LETZTE}", quer=True, titelzeilen="6:6")


def blatt_beschriftung(wb):
    ws = wb.create_sheet("Beschriftung")
    TE_MM = 17.5
    # Excel druckt Spaltenbreiten nicht linear. Kalibriert an der PDF-Ausgabe von Excel:
    # Breite 9,43 -> 18,80 mm, Breite 8,71 -> 17,23 mm  =>  mm = 2,198 x Breite - 1,92
    te_breite = round((TE_MM + 1.92) / 2.198, 3)  # ~ 8,835 fuer 17,5 mm je TE bei 100 %
    breiten(ws, {"A": 1.5, **{get_column_letter(c): te_breite for c in range(2, 14)}, "N": 1.5})
    zelle(ws, "B1", "Beschriftung Verteiler", size=14, bold=True)
    zelle(ws, "B2", '=IF(Anlage!B12="","Verteiler: –","Verteiler: "&Anlage!B12&IF(Anlage!D12="",""," · "&Anlage!D12))',
          size=9, color=GRAU)
    ws.row_dimensions[1].height = 24
    verbinden(ws, "B3:M3", "Mit 100 % drucken (keine Skalierung). Ein Feld = 1 TE = 17,5 mm – am Probedruck nachmessen. "
                           "Inhalt kommt aus der Stromkreisliste (Reihe, Platz, BMK, Beschriftung kurz).",
              size=7, italic=True, color=GRAU, wrap=True, v="top")
    ws.row_dimensions[3].height = 20
    logo(ws, "M", breite_px=96)
    gestrichelt = Side(style="dashed", color=GRAU_HELL)
    schnitt = Border(left=gestrichelt, right=gestrichelt, top=gestrichelt, bottom=gestrichelt)
    bereich_ = f"{SK}!$S${ERSTE}:$S${LETZTE}"
    r = 5
    for reihe in range(1, 9):
        zelle(ws, f"B{r}", f"Reihe {reihe}", size=7, color=GRAU)
        ws.row_dimensions[r].height = 11
        for platz in range(1, 13):
            col = get_column_letter(platz + 1)
            schluessel = f'"{reihe}-{platz}"'
            f = (f"=IFERROR(INDEX({SK}!$B${ERSTE}:$B${LETZTE},MATCH({schluessel},{bereich_},0))&CHAR(10)&"
                 f"INDEX({SK}!$R${ERSTE}:$R${LETZTE},MATCH({schluessel},{bereich_},0)),\"\")")
            c = zelle(ws, f"{col}{r + 1}", f, size=7, h="center", v="center", wrap=True)
            c.border = schnitt
        ws.row_dimensions[r + 1].height = 34  # 12 mm Streifenhoehe
        ws.row_dimensions[r + 2].height = 8
        r += 3
    drucken(ws, f"A1:N{r}", quer=True, skalierung=100)
    ws.print_options.horizontalCentered = False
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.6, header=0.2, footer=0.3)


def blatt_pruefbericht(wb):
    ws = wb.create_sheet("Prüfbericht")
    breiten(ws, {"A": 27, "B": 44, "C": 12, "D": 30})
    titel(ws, "Prüfbericht", "Prüfung elektrischer Anlagen nach DIN VDE 0100-600 bzw. DIN VDE 0105-100", "D")

    def f_anlage(r, lab, formel, lab2=None, formel2=None, span=False, fmt2=None):
        beschriftung(ws, f"A{r}", lab)
        if span:
            formelfeld(ws, f"B{r}:D{r}", formel, merge=True)
        else:
            formelfeld(ws, f"B{r}", formel)
            if lab2:
                beschriftung(ws, f"C{r}", lab2)
                formelfeld(ws, f"D{r}", formel2, fmt=fmt2, h="left")
        ws.row_dimensions[r].height = 18

    a = lambda ref: f'=IF(Anlage!{ref}="","",Anlage!{ref})'
    abschnitt(ws, 5, "Anlage")
    f_anlage(6, "Auftraggeber / Betreiber", a("B7"), "Auftrags-Nr.", a("B9"))
    f_anlage(7, "Anlagenanschrift", a("B8"), span=True)
    f_anlage(8, "Verteiler", a("B12"), "Standort", a("D12"))
    f_anlage(9, "Netzsystem", a("B18"), "Nennspannung", a("D18"))

    abschnitt(ws, 11, "Prüfung")
    beschriftung(ws, "A12", "Prüfart")
    eingabe(ws, "B12:D12", merge=True)
    liste(ws, "B12", ["Erstprüfung nach DIN VDE 0100-600", "Wiederholungsprüfung nach DIN VDE 0105-100",
                      "Prüfung nach Änderung / Erweiterung (DIN VDE 0100-600)"])
    beschriftung(ws, "A13", "Prüfdatum")
    eingabe(ws, "B13", fmt=FMT_DATUM)
    beschriftung(ws, "C13", "Prüfer")
    formelfeld(ws, "D13", a("B26"))
    beschriftung(ws, "A14", "Messgerät (Hersteller, Typ)")
    eingabe(ws, "B14")
    beschriftung(ws, "C14", "Seriennummer")
    eingabe(ws, "D14")
    beschriftung(ws, "A15", "Kalibriert bis")
    eingabe(ws, "B15", fmt=FMT_DATUM)
    beschriftung(ws, "C15", "Erdungswiderst. RA")
    formelfeld(ws, "D15", '=IF(Anlage!D20="","",Anlage!D20&" Ω")')
    for r in range(12, 16):
        ws.row_dimensions[r].height = 18

    def pruefblock(start, nr, titel_, punkte):
        verbinden(ws, f"A{start}:B{start}", f"{nr}  {titel_}", size=8, bold=True, color="FFFFFF", bg=GRAPHIT, rahmen=True)
        zelle(ws, f"C{start}", "Ergebnis", size=8, bold=True, color="FFFFFF", bg=GRAPHIT, h="center", rahmen=True)
        zelle(ws, f"D{start}", "Bemerkung", size=8, bold=True, color="FFFFFF", bg=GRAPHIT, rahmen=True)
        ws.row_dimensions[start].height = 18
        r = start + 1
        for text, formel, bemerkung in punkte:
            verbinden(ws, f"A{r}:B{r}", text, size=9, wrap=True, rahmen=True)
            if formel:
                formelfeld(ws, f"C{r}", formel, h="center")
            else:
                eingabe(ws, f"C{r}", h="center")
                liste(ws, f"C{r}", ["i.O.", "n.i.O.", "entfällt"])
            if bemerkung:
                formelfeld(ws, f"D{r}", bemerkung)
                ws[f"D{r}"].font = font(8, color=GRAU)
            else:
                eingabe(ws, f"D{r}")
            ws.row_dimensions[r].height = 16 if len(text) < 78 else 26
            r += 1
        return r

    r = pruefblock(17, "1", "Besichtigen", [
        ("Auswahl der Betriebsmittel passend zu den äußeren Einflüssen (Schutzart, Brandschutz)", None, None),
        ("Basisschutz: Abdeckungen, Isolierungen und Gehäuse vollständig und unbeschädigt", None, None),
        ("Leiterquerschnitte und Schutzeinrichtungen richtig zugeordnet", None, None),
        ("Kennzeichnung von Stromkreisen, Schutzeinrichtungen und Leitern (N blau, PE grün-gelb)", None, None),
        ("Verteiler zugänglich, Bedienung und Beschriftung in Ordnung", None, None),
        ("Leitungsverlegung, Befestigung und mechanischer Schutz", None, None),
        ("Klemmstellen fest, Leiterverbindungen fachgerecht", None, None),
        ("Schutzpotentialausgleich (Haupt- und zusätzlicher) vorhanden und angeschlossen", None, None),
        ("Brandabschottungen bei Durchführungen", None, None),
        ("Zusätzlicher Schutz durch RCD ≤ 30 mA, wo gefordert (Steckdosen ≤ 32 A, Beleuchtung in Wohnungen, außen)", None, None),
        ("Überspannungsschutz, wo gefordert", None, None),
        ("Dokumentation vorhanden (Stromkreisliste, Pläne)", None, None),
    ])
    r = pruefblock(r, "2", "Erproben", [
        ("Funktion der Fehlerstrom-Schutzeinrichtungen (Prüftaste)", None, None),
        ("Funktion von Schalt-, Steuer- und Verriegelungseinrichtungen", None, None),
        ("Hauptschalter / Not-Aus", None, None),
        ("Funktion von Beleuchtung und Steckdosen (Stichprobe)", None, None),
    ])
    rng = lambda c: f"{MP}!${c}${ERSTE}:${c}${LETZTE}"
    r_messen = r
    r = pruefblock(r, "3", "Messen – Werte siehe Messprotokoll", [
        ("Durchgängigkeit der Schutzleiter und des Potentialausgleichs (RPE)", None, None),
        ("Isolationswiderstand (Riso)",
         f'=IF(COUNT({MP}!$F${ERSTE}:$G${LETZTE})=0,"",IF(SUM({rng("R")})>0,"n.i.O.","i.O."))',
         f'=IF(COUNT({rng("F")})=0,"",COUNT({rng("F")})&" Stromkreise gemessen")'),
        ("Schleifenimpedanz / Abschaltbedingung (Zs)",
         f'=IF(COUNT({rng("I")})=0,"",IF(SUM({rng("S")})>0,"n.i.O.","i.O."))',
         f'=IF(COUNT({rng("I")})=0,"",COUNT({rng("I")})&" Stromkreise gemessen")'),
        ("Fehlerstrom-Schutzeinrichtungen (IΔ, tA, UB)",
         f'=IF(COUNT({MP}!$L${ERSTE}:$N${LETZTE})=0,"",IF(SUM({rng("T")})>0,"n.i.O.","i.O."))',
         f'=IF(COUNT({rng("M")})=0,"",COUNT({rng("M")})&" FI gemessen")'),
        ("Drehfeld der Drehstromsteckdosen",
         f'=IF(COUNTA({rng("O")})=0,"",IF(SUM({rng("U")})>0,"n.i.O.","i.O."))', None),
        ("Erdungswiderstand (bei TT-System)", None, None),
    ])
    ergebnis_farben(ws, f"C18:C{r - 1}")

    abschnitt(ws, r + 1, "Mängel und Hinweise")
    for i in range(4):
        eingabe(ws, f"A{r + 2 + i}:D{r + 2 + i}", merge=True)
        ws.row_dimensions[r + 2 + i].height = 18
    r = r + 7
    abschnitt(ws, r, "Ergebnis")
    beschriftung(ws, f"A{r + 1}", "Gesamtergebnis")
    eingabe(ws, f"B{r + 1}:D{r + 1}", merge=True, wrap=True)
    liste(ws, f"B{r + 1}", ["Keine Mängel – die Anlage entspricht den anerkannten Regeln der Elektrotechnik",
                            "Mängel festgestellt – Beseitigung erforderlich",
                            "Gefährliche Mängel – Anlage(nteil) außer Betrieb genommen"])
    ws.row_dimensions[r + 1].height = 22
    verbinden(ws, f"B{r + 2}:D{r + 2}",
              f'=IF(COUNTIF(C18:C{r - 8},"n.i.O.")>0,"Achtung: mindestens ein Prüfpunkt ist „n.i.O.“ – Gesamtergebnis prüfen.","")',
              size=8, bold=True, color="C00000")
    beschriftung(ws, f"A{r + 3}", "Nächste Prüfung empfohlen")
    formelfeld(ws, f"B{r + 3}", '=IF(B13="","",EDATE(B13,48))', fmt=FMT_DATUM)
    ws[f"B{r + 3}"].comment = Comment("Prüfdatum + 4 Jahre (Empfehlung für Wohngebäude). Bei Bedarf überschreiben.",
                                      _autor())
    beschriftung(ws, f"C{r + 3}", "Prüfplakette")
    eingabe(ws, f"D{r + 3}")
    liste(ws, f"D{r + 3}", ["angebracht", "nicht angebracht"])
    ws.row_dimensions[r + 3].height = 18
    # Unterschriften
    ws.row_dimensions[r + 4].height = 30
    ws.row_dimensions[r + 5].height = 4
    for col in "AB":
        ws[f"{col}{r + 5}"].border = Border(bottom=STRICH)
    ws[f"D{r + 5}"].border = Border(bottom=STRICH)
    zelle(ws, f"A{r + 6}", "Ort, Datum", size=8, color=GRAU)
    zelle(ws, f"D{r + 6}", "Unterschrift Prüfer (Elektrofachkraft)", size=8, color=GRAU)
    drucken(ws, f"A1:D{r + 6}", eine_seite=True)
    return r_messen


def blatt_messprotokoll(wb):
    ws = wb.create_sheet(MP)
    spalten = {
        "A": ("Nr.", 4.5), "B": ("BMK", 7), "C": ("Stromkreis / Verbraucher", 28), "D": ("Schutzorgan", 13),
        "E": ("Leitung", 12), "F": ("Riso L–PE (MΩ)", 8), "G": ("Riso N–PE (MΩ)", 8), "H": ("RPE (Ω)", 7),
        "I": ("Zs L–PE (Ω)", 7.5), "J": ("Ik (A)", 7), "K": ("Zs max (Ω)", 7.5), "L": ("IΔ (mA)", 7),
        "M": ("tA (ms)", 7), "N": ("UB (V)", 6.5), "O": ("Dreh-feld", 7.5), "P": ("Ergebnis", 8.5),
        "Q": ("Bemerkung", 20),
    }
    breiten(ws, {k: v[1] for k, v in spalten.items()})
    breiten(ws, {"R": 4, "S": 4, "T": 4, "U": 4})
    titel(ws, "Messprotokoll", f"Messwerte je Stromkreis · {FIRMA.get('name', '')}", "Q")

    beschriftung(ws, "A4", "Verteiler")
    ws.merge_cells("A4:B4")
    formelfeld(ws, "C4", '=IF(Anlage!B12="","",Anlage!B12&IF(Anlage!D12="",""," · "&Anlage!D12))', bold=True)
    beschriftung(ws, "D4", "Objekt")
    formelfeld(ws, "E4:J4", '=IF(Anlage!B8="","",Anlage!B8)', merge=True)
    beschriftung(ws, "K4", "Datum")
    ws.merge_cells("K4:L4")
    formelfeld(ws, "M4:N4", f'=IF({PB}!B13="","",{PB}!B13)', merge=True, fmt=FMT_DATUM, h="center")
    beschriftung(ws, "O4", "Gerät")
    formelfeld(ws, "P4:Q4", f'=IF({PB}!B14="","",{PB}!B14)', merge=True)
    ws.row_dimensions[4].height = 18

    # Grenzwerte (Eingabe, aenderbar)
    beschriftung(ws, "A5", "Grenzwerte")
    ws.merge_cells("A5:B5")
    grenz = [("C5", "Nennspannung U0 (V)", "D5", 230), ("E5", "Riso min (MΩ)", "F5", 1),
             ("G5", "Faktor Zs", "H5", "=2/3"), ("I5", "tA max (ms)", "J5", 300), ("K5", "UB max (V)", "L5", 50)]
    for lab_ref, lab, ref, w in grenz:
        zelle(ws, lab_ref, lab, size=8, color=GRAU, h="right", rahmen=True, bg=FLAECHE, wrap=True)
        eingabe(ws, ref, wert=w, h="center", fmt="0.00" if ref == "H5" else None)
    verbinden(ws, "M5:Q5", "Zs max = Faktor × U0 / Ia  ·  Ia: B = 5·In, C = 10·In, D = 20·In, K = 14·In, Z = 3·In. "
                           "Bewertung ist eine Hilfe – maßgeblich sind Norm und Herstellerangaben.",
              size=7, italic=True, color=GRAU, wrap=True)
    ws.row_dimensions[5].height = 26
    ws["H5"].comment = Comment("Faktor 2/3 berücksichtigt die Erwärmung der Leiter im Betrieb "
                               "(Bewertung gemessener Werte nach DIN VDE 0100-600). Für reine Rechenwerte 1 eintragen.",
                               _autor())
    ws["F5"].comment = Comment("Mindestwert Isolationswiderstand bei Prüfspannung 500 V DC: 1 MΩ.", _autor())
    ws["J5"].comment = Comment("Abschaltzeit bei IΔn für unverzögerte RCD: höchstens 300 ms.", _autor())

    kopfzeile(ws, 6, [(k, v[0]) for k, v in spalten.items()], hoehe=34)
    for col, t in zip("RSTU", ["Riso", "Zs", "FI", "Dreh"]):
        zelle(ws, f"{col}6", t, size=6, color=GRAU_HELL)

    s = SK
    for r in range(ERSTE, LETZTE + 1):
        b = f"{s}!B{r}"
        zelle(ws, f"A{r}", f'=IF({b}="","",ROW()-{ERSTE - 1})', rahmen=True, h="center", size=9, color=GRAU)
        zelle(ws, f"B{r}", f'=IF({b}="","",{b})', rahmen=True, size=9, bold=True)
        zelle(ws, f"C{r}", f'=IF({b}="","",{s}!P{r}&IF(AND({s}!P{r}<>"",{s}!Q{r}<>"")," – ","")&{s}!Q{r})',
              rahmen=True, size=8)
        zelle(ws, f"D{r}",
              f'=IF({b}="","",IF({s}!F{r}="FI/RCD","FI "&{s}!H{r}&" A / "&{s}!K{r}&" mA",'
              f'IF({s}!F{r}="FI/LS (RCBO)","FI/LS "&{s}!G{r}&{s}!H{r}&" / "&{s}!K{r}&" mA",'
              f'IF({s}!G{r}="",{s}!F{r}&IF({s}!H{r}="",""," "&{s}!H{r}&" A"),{s}!G{r}&{s}!H{r}))))',
              rahmen=True, size=8)
        zelle(ws, f"E{r}", f'=IF({b}="","",TRIM({s}!M{r}&" "&{s}!N{r}))', rahmen=True, size=8)
        for col in "FGHILMNOQ":
            zelle(ws, f"{col}{r}", bg=EINGABE, rahmen=True, size=9, h="left" if col == "Q" else "center",
                  fmt="0.00" if col in "FGHI" else None)
        zelle(ws, f"J{r}", f'=IF(N(I{r})>0,$D$5/I{r},"")', rahmen=True, size=9, h="center", fmt="0")
        zelle(ws, f"K{r}",
              f'=IF(OR({b}="",N({s}!H{r})=0),"",IFERROR($H$5*$D$5/({s}!H{r}*INDEX({{5,10,20,14,3}},'
              f'MATCH({s}!G{r},{{"B","C","D","K","Z"}},0))),""))',
              rahmen=True, size=9, h="center", fmt="0.00")
        zelle(ws, f"P{r}",
              f'=IF(COUNTA(F{r}:I{r},L{r}:O{r})=0,"",IF(R{r}+S{r}+T{r}+U{r}>0,"n.i.O.","i.O."))',
              rahmen=True, size=9, h="center", bold=True)
        # Hilfsspalten (ausgeblendet): 1 = Grenzwert verletzt
        ws[f"R{r}"] = f"=IF(OR(AND(ISNUMBER(F{r}),F{r}<$F$5),AND(ISNUMBER(G{r}),G{r}<$F$5)),1,0)"
        ws[f"S{r}"] = f"=IF(AND(ISNUMBER(I{r}),N(K{r})>0,I{r}>N(K{r})),1,0)"
        ws[f"T{r}"] = (f"=IF(OR(AND(ISNUMBER(L{r}),N({s}!K{r})>0,OR(L{r}>N({s}!K{r}),L{r}<0.5*N({s}!K{r}))),"
                       f"AND(ISNUMBER(M{r}),M{r}>$J$5),AND(ISNUMBER(N{r}),N{r}>$L$5)),1,0)")
        ws[f"U{r}"] = f'=IF(O{r}="links",1,0)'
        for col in "RSTU":
            ws[f"{col}{r}"].font = font(7, color=GRAU_HELL)
        ws.row_dimensions[r].height = 16

    rng = lambda c: f"{c}{ERSTE}:{c}{LETZTE}"
    liste(ws, rng("O"), ["rechts", "links", "–"])
    ergebnis_farben(ws, rng("P"))
    rot = fill("F8CBAD")
    ws.conditional_formatting.add(f"F{ERSTE}:G{LETZTE}", FormulaRule(formula=[f"$R{ERSTE}=1"], fill=rot))
    ws.conditional_formatting.add(rng("I"), FormulaRule(formula=[f"$S{ERSTE}=1"], fill=rot))
    ws.conditional_formatting.add(f"L{ERSTE}:N{LETZTE}", FormulaRule(formula=[f"$T{ERSTE}=1"], fill=rot))
    for col in "RSTU":
        ws.column_dimensions[col].hidden = True

    # Unterschrift
    u = LETZTE + 2
    ws.row_dimensions[u].height = 28
    for col in "BC":
        ws[f"{col}{u}"].border = Border(bottom=STRICH)
    for col in "IJKLMNOP":
        ws[f"{col}{u}"].border = Border(bottom=STRICH)
    zelle(ws, f"B{u + 1}", "Ort, Datum", size=8, color=GRAU)
    zelle(ws, f"I{u + 1}", '="Unterschrift Prüfer: "&Anlage!B26', size=8, color=GRAU)
    ws.auto_filter.ref = f"A6:Q{LETZTE}"
    ws.freeze_panes = "D7"
    drucken(ws, f"A1:Q{u + 1}", quer=True, titelzeilen="6:6")


VORLAGE_NAME = "Verteilerdokumentation Stromkreisliste Messprotokoll.xltx"


def _verbinde(*teile) -> str:
    return " · ".join(str(x).strip() for x in teile if x and str(x).strip())


def _person() -> str:
    return ", ".join(x for x in (FIRMA.get("inhaber"), FIRMA.get("titel")) if x)


def _autor() -> str:
    return FIRMA.get("inhaber") or FIRMA.get("name") or "elektroplan"


def erzeuge(ziel_ordner: Path, firma: dict | None = None, logo_pfad: Path | None = None) -> Path:
    """Schreibt die Excel-Vorlage (.xltx) mit Firmendaten und Logo nach ziel_ordner."""
    global FIRMA, LOGO_XL, LOGO_RATIO, KUPFER, KUPFER_LINIE
    FIRMA = dict(firma if firma is not None else _firma.lade())
    if logo_pfad is None and firma is None:
        logo_pfad = _firma.logo_pfad()
    farbe = str(FIRMA.get("akzentfarbe") or "").lstrip("#").upper()
    if len(farbe) == 6:
        KUPFER = farbe
        KUPFER_LINIE = Side(style="medium", color=KUPFER)
    ziel_ordner = Path(ziel_ordner)
    ziel_ordner.mkdir(parents=True, exist_ok=True)
    LOGO_XL = None
    if logo_pfad and Path(logo_pfad).exists():
        bild = PILImage.open(logo_pfad).convert("RGBA")
        LOGO_RATIO = bild.width / bild.height
        LOGO_XL = ziel_ordner / "_logo_excel.png"
        bild.resize((520, max(1, round(520 / LOGO_RATIO))), PILImage.LANCZOS).save(LOGO_XL)

    wb = Workbook()
    wb.properties.creator = _autor()
    blatt_anleitung(wb)
    blatt_anlage(wb)
    blatt_stromkreisliste(wb)
    blatt_beschriftung(wb)
    blatt_pruefbericht(wb)
    blatt_messprotokoll(wb)
    wb.active = 1
    wb.calculation.fullCalcOnLoad = True      # Excel rechnet beim Öffnen alle Formeln neu
    wb.template = True
    ziel = ziel_ordner / VORLAGE_NAME
    wb.save(ziel)
    if LOGO_XL is not None:
        LOGO_XL.unlink(missing_ok=True)
    return ziel
