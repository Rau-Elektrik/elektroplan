"""Verteilerplan-Daten aus der Excel-Vorlage „Verteilerdokumentation Stromkreisliste
Messprotokoll" (Blätter „Anlage" und „Stromkreisliste") lesen."""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path


def _wert(ws, zelle):
    v = ws[zelle].value
    if v is None:
        return ""
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _leitung_teilen(text: str):
    """"NYM-J 5×10 mm²" → ("NYM-J", "5×10")."""
    m = re.match(r"\s*([A-Za-z0-9\-/]+)\s+(\d+\s*[x×*]\s*[\d,\.]+)", text or "")
    if m:
        return m.group(1), m.group(2).replace(" ", "")
    return text, ""


def lese_stromkreisliste(pfad: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(pfad, data_only=True)
    if "Stromkreisliste" not in wb.sheetnames:
        raise ValueError(f"{pfad.name}: Blatt „Stromkreisliste“ fehlt – ist das die "
                         "Verteilerdokumentations-Vorlage?")
    an = wb["Anlage"] if "Anlage" in wb.sheetnames else None
    sk = wb["Stromkreisliste"]

    # Kopfzeile suchen (Spalte B = "BMK")
    kopf = next((r for r in range(1, 20) if str(sk[f"B{r}"].value or "").strip() == "BMK"), 6)
    spalten = {}
    for zelle in sk[kopf]:
        if zelle.value:
            spalten[str(zelle.value).strip().lower()] = zelle.column_letter

    def sp(*namen):
        for n in namen:
            for k, v in spalten.items():
                if k.startswith(n):
                    return v
        return None

    S = dict(bmk=sp("bmk"), art=sp("schutzorgan"), char=sp("char"), In=sp("in (a)", "in"),
             pole=sp("pole"), fi=sp("vor-geschalteter fi", "vorgeschalteter fi", "vor"),
             idn=sp("iδn", "idn", "i∆n"), fityp=sp("fi-typ"), leitung=sp("leitung"),
             adern=sp("adern"), laenge=sp("länge", "laenge"), raum=sp("raum"),
             verbraucher=sp("verbraucher"))
    eintraege = []
    for r in range(kopf + 1, sk.max_row + 1):
        bmk = _wert(sk, f"{S['bmk']}{r}")
        if not bmk:
            continue
        g = lambda k: _wert(sk, f"{S[k]}{r}") if S.get(k) else ""
        e = {"bmk": bmk, "art": g("art"), "char": g("char"), "in": g("In"), "pole": g("pole"),
             "fi": g("fi"), "idn": g("idn"), "fi_typ": g("fityp"), "leitung": g("leitung"),
             "adern": g("adern"), "laenge": g("laenge"), "raum": g("raum"),
             "verbraucher": g("verbraucher")}
        eintraege.append({k: v for k, v in e.items() if v})

    daten = {"plan": "verteilerplan", "schutzorgane": eintraege,
             "dateiname": f"Verteilerplan {Path(pfad).stem}"}
    if an is not None:
        zul_l, zul_a = _leitung_teilen(_wert(an, "B14"))
        daten["projekt"] = {"kunde": _wert(an, "B7"), "objekt": _wert(an, "B8"),
                            "auftrag": _wert(an, "B9"), "datum": _wert(an, "D9")}
        daten["verteiler"] = {"name": _wert(an, "B12"), "standort": _wert(an, "D12"),
                              "netz": _wert(an, "B18")}
        daten["einspeisung"] = {"leitung": zul_l, "adern": zul_a,
                                "vorsicherung": _wert(an, "D14")}
    return daten


# ------------------------------------------------------------------ Schreiben (Gegenstück)

from .excel_vorlage import VORLAGE_NAME  # noqa: E402
ERSTE_ZEILE, LETZTE_ZEILE = 7, 56
ART_VORLAGE = {"LS": "LS", "FI": "FI/RCD", "FILS": "FI/LS (RCBO)", "HS": "Hauptschalter", "SLS": "SLS",
               "SPD": "Überspannungsschutz", "SICH": "Neozed"}
KURZ = [("Steckdosen", "SD"), ("Steckdose", "SD"), ("Beleuchtung", "Li"), ("Licht", "Li"),
        ("Waschmaschine", "WM"), ("Wäschetrockner", "WT"), ("Trockner", "WT"), ("Geschirrspüler", "GS"),
        ("Durchlauferhitzer", "DLE"), ("Kühlschrank", "KS"), ("Herd / Kochfeld", "Herd"), ("Kochfeld", "Herd"),
        ("Außensteckdose", "Außen SD"), ("Arbeitsplatte", "AP"), ("Schlafzimmer", "Schlafzi."),
        ("Wohnzimmer", "Wohnzi."), ("Kinderzimmer", "Kinderzi.")]


def vorlage_finden() -> Path | None:
    """Eigene Excel-Vorlage aus dem Konfigurationsordner (siehe ep/firma.py)."""
    from .firma import stromkreisliste_vorlage
    return stromkreisliste_vorlage()


def _te(e: dict, art: str) -> int:
    pole = str(e.get("pole") or "").replace(" ", "")
    if art in ("FI", "SPD"):
        return 2 if pole in ("2", "1+N") else 4
    if art == "FILS":
        return 2 if pole in ("", "1", "1+N", "2") else 4
    return {"1": 1, "1+N": 2, "2": 2, "3": 3, "3+N": 4, "4": 4}.get(pole, 1)


def _kurz(e: dict, art: str) -> str:
    def zahl(v):
        s = str(v or "").strip()
        return s[:-2] if s.endswith(".0") else s
    if art == "FI":
        return f"FI {zahl(e.get('in'))}A {zahl(e.get('idn')) or 30}mA"[:14]
    if art == "HS":
        return "Hauptschalter"
    if art == "SPD":
        return "Überspannung"
    t = str(e.get("verbraucher") or e.get("raum") or "").strip()
    raum = str(e.get("raum") or "").strip()
    for lang, kurz in KURZ:
        t = t.replace(lang, kurz)
    if raum and raum.lower() not in t.lower():
        r = raum
        for lang, kurz in KURZ:
            r = r.replace(lang, kurz)
        t = f"{r} {t}".strip()
    return t[:14]


def _datum(s: str):
    try:
        return dt.datetime.strptime(s.strip(), "%d.%m.%Y")
    except (ValueError, AttributeError):
        return s or None


def schreibe_stromkreisliste(daten: dict, ziel: Path, vorlage: Path | None = None) -> tuple[Path | None, list[str]]:
    """Füllt die eigene Excel-Vorlage „Verteilerdokumentation Stromkreisliste Messprotokoll“ aus der Plan-JSON.

    Blatt „Anlage“ (Kopf-, Verteiler-, Netzdaten) und „Stromkreisliste“ (eine Zeile je Schutzorgan, Reihe/Platz
    fortlaufend mit `verteiler.te_je_reihe`, Standard 12 TE). Messprotokoll und Prüfbericht bleiben leer.
    Rückgabe: (Pfad oder None, Hinweise)."""
    import openpyxl
    hinweise: list[str] = []
    vorlage = vorlage or vorlage_finden()
    if not vorlage:
        return None, [f"Stromkreisliste: Excel-Vorlage „{VORLAGE_NAME}“ nicht gefunden – keine Excel-Datei "
                      "erzeugt. Abhilfe: python einrichtung.py vorlage"]
    from .verteilerplan import _art
    wb = openpyxl.load_workbook(vorlage)
    wb.template = False
    proj = daten.get("projekt") or {}
    vt = daten.get("verteiler") or {}
    ein = daten.get("einspeisung") or {}
    an = wb["Anlage"]
    an["B7"] = proj.get("kunde") or None
    an["B8"] = proj.get("objekt") or None
    an["B9"] = proj.get("auftrag") or None
    an["D9"] = _datum(proj.get("datum") or "")
    an["B12"] = vt.get("name") or None
    an["D12"] = vt.get("standort") or None
    an["B13"] = vt.get("hersteller") or None
    if vt.get("reihen") and vt.get("te_je_reihe"):
        an["D13"] = f"{vt['reihen']} × {vt['te_je_reihe']}"
    zul = " ".join(str(x) for x in (ein.get("leitung"), ein.get("adern")) if x)
    an["B14"] = (zul + (" mm²" if ein.get("adern") else "")) or None
    an["D14"] = ein.get("vorsicherung") or None
    if vt.get("netz") in ("TN-C-S", "TN-S", "TT", "TN-C"):
        an["B18"] = vt["netz"]
    spd = next((e for e in daten.get("schutzorgane") or [] if _art(e) == "SPD"), None)
    if spd:
        typ = str(spd.get("typ") or "").replace("Typ", "").strip()
        an["B20"] = f"Typ {typ}" if f"Typ {typ}" in ("Typ 1+2", "Typ 2", "Typ 3") else None

    sk = wb["Stromkreisliste"]
    for r in range(ERSTE_ZEILE, LETZTE_ZEILE + 1):          # Beispielzeilen der Vorlage leeren
        for c in range(2, 19):
            if not str(sk.cell(r, c).value or "").startswith("="):
                sk.cell(r, c).value = None
    te_reihe = int(vt.get("te_je_reihe") or 12)
    reihe, platz = 1, 1
    eintraege = daten.get("schutzorgane") or []
    if len(eintraege) > LETZTE_ZEILE - ERSTE_ZEILE + 1:
        hinweise.append(f"Stromkreisliste: {len(eintraege)} Schutzorgane, die Vorlage hat nur "
                        f"{LETZTE_ZEILE - ERSTE_ZEILE + 1} Zeilen – Rest fehlt")
    for i, e in enumerate(eintraege[:LETZTE_ZEILE - ERSTE_ZEILE + 1]):
        r = ERSTE_ZEILE + i
        art = _art(e)
        te = _te(e, art)
        if art != "ABZ":
            if platz + te - 1 > te_reihe:
                reihe, platz = reihe + 1, 1
            sk.cell(r, 3).value, sk.cell(r, 4).value, sk.cell(r, 5).value = reihe, platz, te
            platz += te
        def zahl(v):
            try:
                f = float(str(v).replace(",", "."))
                return int(f) if f.is_integer() else f
            except (TypeError, ValueError):
                return v or None
        werte = {2: e.get("bmk"), 6: ART_VORLAGE.get(art, "Sonstiges"), 7: e.get("char") or None,
                 8: zahl(e.get("in")), 9: str(e.get("pole")) if e.get("pole") not in (None, "") else None,
                 10: e.get("fi") or None, 11: zahl(e.get("idn")), 12: e.get("fi_typ") or None,
                 13: e.get("leitung") or None, 14: (str(e.get("adern") or "").replace("x", "×") or None),
                 15: zahl(e.get("laenge")), 16: e.get("raum") or None,
                 17: e.get("verbraucher") or e.get("hinweis") or None, 18: _kurz(e, art) or None}
        if art == "ABZ":
            werte[6] = "Sonstiges"
        for c, v in werte.items():
            sk.cell(r, c).value = v
    if any(sk.cell(r, 3).value and sk.cell(r, 3).value > 1 for r in range(ERSTE_ZEILE, LETZTE_ZEILE + 1)) \
            and not vt.get("te_je_reihe"):
        hinweise.append(f"Stromkreisliste: Reihe/Platz mit {te_reihe} TE je Reihe angenommen – "
                        "bei anderem Verteiler `verteiler.te_je_reihe` angeben")
    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)
    return ziel, hinweise
